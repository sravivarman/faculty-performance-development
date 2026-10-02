"""Read a journal workbook without editing it or importing any publications.

Run with Python containing openpyxl (the bundled workspace runtime supports it).
"""
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def extract(path, sheet_name=None):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        candidates = []
        for sheet in workbook:
            if sheet_name and sheet.title != sheet_name:
                continue
            for row in sheet.iter_rows(min_row=1, max_row=min(30, sheet.max_row)):
                headers = [str(c.value or "").strip().casefold() for c in row]
                if "authors" in headers and "title of the paper" in headers:
                    candidates.append((sheet, row[0].row, headers.index("authors"), headers.index("title of the paper")))
                    break
        if len(candidates) != 1:
            raise ValueError("Select one journal sheet with Authors and Title of the paper headers using --sheet")
        sheet, header_row, author_index, title_index = candidates[0]
        rows = []
        for row in sheet.iter_rows(min_row=header_row+1):
            # Formatting-only rows are ignored; every nonempty data row is retained,
            # including missing titles/claimants, for explicit reconciliation review.
            if not any(c.value is not None and str(c.value).strip() for c in row):
                continue
            if any(c.data_type == "f" for c in (row[author_index], row[title_index])):
                raise ValueError(f"Uncalculated formula in claimant/title at Excel row {row[0].row}; review the workbook first")
            rows.append({"excel_row": row[0].row, "author_value": str(row[author_index].value or ""),
                         "has_title": bool(str(row[title_index].value or "").strip())})
        return {"source": str(Path(path).resolve()), "source_sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest(),
                "sheet": sheet.title, "header_row": header_row, "rows": rows}
    finally:
        workbook.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("excel")
    parser.add_argument("--sheet")
    args = parser.parse_args()
    print(json.dumps(extract(args.excel, args.sheet), ensure_ascii=True))
