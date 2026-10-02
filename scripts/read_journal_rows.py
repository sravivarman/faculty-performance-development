"""Read historical journal inputs, without modifying the workbook."""
import hashlib
import json
import sys
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook


def extract(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book['Journal']
        expected = ['S.No', 'Title of the paper', 'Authors', 'Name of the Journal', 'DoI',
                    'Month and Year', 'Indexing', 'Impact Factor', 'Quartile']
        headers = [str(c.value or '').strip() for c in next(sheet.iter_rows())][:9]
        if headers != expected:
            raise ValueError(f'Unexpected headers: {headers}')
        rows = []
        for cells in sheet.iter_rows(min_row=2):
            if not any(c.value is not None and str(c.value).strip() for c in cells):
                continue
            if any(c.data_type == 'f' for c in cells[:9]):
                raise ValueError(f'Formula requires review at row {cells[0].row}')
            values = [c.value.isoformat() if isinstance(c.value, (date, datetime)) else c.value for c in cells[:9]]
            rows.append(dict(excel_row=cells[0].row, **dict(zip(expected, values))))
        return {'source': str(Path(path).resolve()), 'sheet': 'Journal',
                'source_sha256': hashlib.sha256(Path(path).read_bytes()).hexdigest(), 'rows': rows}
    finally:
        book.close()

if __name__ == '__main__':
    print(json.dumps(extract(sys.argv[1]), ensure_ascii=True))
